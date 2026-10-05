#!/usr/bin/env python3
"""
Хабаровск — полный сбор и анализ клиник под вывод инъекционного препарата
Источники: 2GIS (публичный API) + Яндекс.Карты (Playwright)
Пайплайн: GeoHarvest → SmartCascadeDedup → SmartCrawler → HybridAnalyzer (Regex+Gemini) → Excel
Поддержка:
 - Умная каскадная склейка 2GIS + Яндекс (по доменам, телефонам, бренду и улице)
 - Обход российских SSL (Минцифры / НУЦ РФ) с verify=False
 - Поддержка русских доменов (.рф / Punycode IDN)
 - Поддержка соцсетей и визиток (VK, Taplink, T.me)
 - Защита от бот-блокировок (DDoS-Guard / Cloudflare / 403)
 - Фоллбек на сниппет и рубрики 2GIS/Яндекс при недоступности сайта
 - Multi-Key Gemini Pool с ротацией и сохранением чекпоинтов
"""

import asyncio
import json
import logging
import os
import re
import signal
import sys
import time
from dataclasses import dataclass, field, asdict
from pathlib import Path
from typing import List, Optional, Dict, Any, Tuple
from urllib.parse import urlparse, urljoin, quote

import httpx
import pandas as pd
from bs4 import BeautifulSoup
from diskcache import Cache
from dotenv import load_dotenv
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field

# Загрузка переменных окружения из .env файла
load_dotenv()

# ====================== ПУТИ И ДИРЕКТОРИИ ======================
DATA_DIR = Path("data")
DATA_DIR.mkdir(exist_ok=True)
RAW_FILE = DATA_DIR / "raw_clinics.json"
CRAWLED_FILE = DATA_DIR / "crawled_sites.json"
CHECKPOINT_FILE = DATA_DIR / "analyzed_checkpoint.json"
OUTPUT_EXCEL = "output/Хабаровск_Клиники_Пилот.xlsx"

# ====================== НАСТРОЙКИ ======================
_raw_keys = os.getenv("GEMINI_API_KEYS", "") or os.getenv("GEMINI_API_KEY", "")
GEMINI_KEYS = [k.strip() for k in _raw_keys.split(",") if k.strip() and not k.strip().startswith("ВСТАВЬ")]

_env_2gis_keys = os.getenv("PUBLIC_KEYS_2GIS", "")
if _env_2gis_keys:
    PUBLIC_KEYS_2GIS = [k.strip() for k in _env_2gis_keys.split(",") if k.strip()]
else:
    PUBLIC_KEYS_2GIS = ["ruzhub8994", "rurbbn3446", "ruregt3044"]

USE_2GIS = os.getenv("USE_2GIS", "true").lower() in ("true", "1", "yes")
USE_YANDEX = os.getenv("USE_YANDEX", "true").lower() in ("true", "1", "yes")

CONFIG = {
    "city": os.getenv("CITY", "Хабаровск"),
    "search_queries": [
        "косметология",
        "медицинский центр",
        "пародонтология",
        "стоматологическая хирургия",
        "челюстно-лицевая хирургия",
        "ортодонтия",
        "ортопедия",
        "травматология",
        "неврология",
        "гинекология",
    ],
    "stop_words_gov": [
        "гбуз", "мбуз", "гбузкк", "военн", "госпиталь",
        "поликлиника №", "краевая клиническая", "городская больница",
        "краевой", "областн", "бюджетн", "ведомствен"
    ],
    "prp_brands": {
        "Cortexil": ["cortexil", "кортексил"],
        "RegenLab": ["regen lab", "regenlab", "редженлаб", "регенлаб"],
        "Plasmolifting": ["plasmolifting", "плазмолифтинг"],
        "Endoret": ["endoret", "эндорет", "prgf"],
        "Medical Case": ["medical case", "медикал кейс"],
        "Ycellbio": ["ycellbio", "вайселбио"],
        "T-Lab": ["t-lab", "т-лаб", "т лаб"],
    },
    "crawler": {
        "concurrency": int(os.getenv("CRAWLER_CONCURRENCY", "8")),
        "timeout": int(os.getenv("CRAWLER_TIMEOUT", "15")),
        "max_pages_per_site": int(os.getenv("CRAWLER_MAX_PAGES", "4")),
        "use_playwright_fallback": os.getenv("USE_PLAYWRIGHT_FALLBACK", "true").lower() in ("true", "1", "yes"),
    },
    "gemini": {
        "model": os.getenv("GEMINI_MODEL", "gemini-3.5-flash-lite"),
        "temperature": float(os.getenv("GEMINI_TEMPERATURE", "0.1")),
        "rpm": int(os.getenv("GEMINI_RPM", "14")),
        "tpm": int(os.getenv("GEMINI_TPM", "250000")),
    },
    "max_pages_2gis": int(os.getenv("MAX_PAGES_2GIS", "5")),
    "max_results_yandex_per_query": int(os.getenv("MAX_RESULTS_YANDEX", "70")),
}

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s | %(levelname)-7s | %(message)s",
    datefmt="%H:%M:%S"
)
logger = logging.getLogger("habarovsk")
cache = Cache("./cache")

# ====================== МОДЕЛИ ======================
class ClinicRecord(BaseModel):
    index: int = Field(alias="№ п/п")
    name: str = Field(alias="Название")
    branches_count: int = Field(default=1, alias="кол-во филиалов")
    activity: str = Field(alias="Деятельность")
    phone: str = Field(alias="Контакт номер")
    email: str = Field(default="", alias="Email")
    site: str = Field(default="", alias="Сайт")
    specialists_count: str = Field(default="Требует уточнения", alias="Кол-во спецов")
    injection_therapy: str = Field(default="Нет", alias="Инъекционная терапия")
    prp_therapy_brand: str = Field(default="Нет", alias="Плазмотерапия (бренд)")
    status: str = Field(default="ok", alias="Статус сбора")
    note: str = Field(default="", alias="Примечание")
    source: str = Field(default="", alias="Источник")
    source_id: str = Field(default="", alias="ID")
    model_config = {"populate_by_name": True}


@dataclass
class RawClinic:
    source_id: str
    name: str
    address: str = ""
    phones: List[str] = field(default_factory=list)
    site: str = ""
    rubrics: List[str] = field(default_factory=list)
    branches_count: int = 1
    source: str = "2gis"
    snippet: str = ""


# ====================== ПРОМПТ GEMINI ======================
GEMINI_PROMPT = """
Ты — эксперт по аудиту коммерческих медицинских центров и клиник России.
Твоя задача — определить, проводится ли в данном учреждении ИНЪЕКЦИОННАЯ ТЕРАПИЯ, и выявить применение ПЛАЗМОТЕРАПИИ (PRP).

Верни СТРОГО валидный JSON без markdown и лишнего текста:
{{
  "injection_therapy": "Да" | "Нет" | "Неясно",
  "injection_reason": "краткое перечисление найденных процедур или обоснование 'Нет'",
  "prp_therapy": "Да" | "Нет" | "Неясно",
  "prp_brands": ["бренд1", "бренд2"],
  "prp_brand_raw": "как написано на сайте или null",
  "specialists_count": 5 | "Требует уточнения",
  "specialists_comment": "пояснение по составу врачей",
  "confidence": 0.85
}}

ПРАВИЛА ОПРЕДЕЛЕНИЯ ИНЪЕКЦИОННОЙ ТЕРАПИИ:
1. КОСМЕТОЛОГИЯ:
   - "Да": биоревитализация, мезотерапия, ботулинотерапия (ботокс, диспорт, ксеомин, релатокс, миотокс), контурная пластика филлерами (губы, скулы, носогубки), мезонити, нитевой лифтинг, коллагенотерапия (Сферогель, Коллост, Nithya), полимолочная кислота (AestheFill, Sculptra, MiraLine).
2. СТОМАТОЛОГИЯ:
   - "Да": плазмолифтинг / PRP десен (лечение пародонтита, гингивита), инъекции гиалуроновой кислоты в десневые сосочки, аутоплазма (PRGF/Endoret) при имплантации и синус-лифтинге, ботулинотерапия при бруксизме и дисфункции ВНЧС.
   - "Нет": обычное лечение кариеса, пломбирование, профгигиена, отбеливание, брекеты. РУТИННАЯ МЕСТНАЯ АНЕСТЕЗИЯ (укол перед сверлением зуба) НЕ СЧИТАЕТСЯ инъекционной терапией! Если в стоматологии нет инъекций в десны/плазмолифтинга/ботокса — ставь "Нет".
3. ОРТОПЕДИЯ, ТРАВМАТОЛОГИЯ, НЕВРОЛОГИЯ:
   - "Да": внутрисуставные инъекции (гиалуроновая кислота, заменители синовиальной жидкости), PRP/PRGF суставов и позвоночника, лечебные медикаментозные блокады, ботулинотерапия при спастике/мигрени.
4. ГИНЕКОЛОГИЯ:
   - "Да": интимная контурная пластика (филлеры гиалуроновой кислоты), интимная биоревитализация, плазмолифтинг / PRP в гинекологии.

ПРАВИЛА ПО ПЛАЗМОТЕРАПИИ:
- Плазмолифтинг, PRP, PRGF, аутоплазмотерапия, плазмотерапия.
- Извлекай бренды только если они явно указаны: Cortexil, RegenLab, Plasmolifting, Endoret, Medical Case, Ycellbio, T-Lab.

ПРАВИЛА ПО СПЕЦИАЛИСТАМ:
- Считай только профильных врачей (косметологи, дерматовенерологи, стоматологи-хирурги/пародонтологи/ортопеды, ортопеды-травматологи, неврологи, гинекологи).
- Если перечня врачей на сайте нет — пиши "Требует уточнения".

Текст сайта или карточки организации:
---
{site_text}
---
"""

# ====================== GEMINI RATE LIMITER & POOL ======================
class KeyRateLimiter:
    """Ограничитель частоты запросов для конкретного ключа Gemini."""
    def __init__(self, key: str, rpm: int = 14, tpm: int = 250000, min_interval: float = 4.2):
        self.key = key
        self.rpm = rpm
        self.tpm = tpm
        self.min_interval = min_interval
        self.last_call_time = 0.0
        self.lock = asyncio.Lock()
        self.history: List[List[float]] = []

    def _purge(self, now: float):
        cutoff = now - 60.0
        self.history = [rec for rec in self.history if rec[0] > cutoff]

    async def acquire(self, estimated_tokens: int) -> int:
        async with self.lock:
            while True:
                now = time.monotonic()
                self._purge(now)

                wait_interval = max(0.0, self.min_interval - (now - self.last_call_time))

                wait_rpm = 0.0
                if len(self.history) >= self.rpm:
                    oldest_ts = self.history[0][0]
                    wait_rpm = max(0.0, 60.0 - (now - oldest_ts) + 0.1)

                wait_tpm = 0.0
                current_tokens = sum(rec[1] for rec in self.history)
                if current_tokens + estimated_tokens > self.tpm:
                    needed_to_free = (current_tokens + estimated_tokens) - self.tpm
                    freed = 0.0
                    for ts, tok in self.history:
                        freed += tok
                        if freed >= needed_to_free:
                            wait_tpm = max(0.0, 60.0 - (now - ts) + 0.1)
                            break

                wait_time = max(wait_interval, wait_rpm, wait_tpm)
                if wait_time > 0:
                    await asyncio.sleep(wait_time)
                else:
                    self.last_call_time = time.monotonic()
                    rec = [self.last_call_time, float(estimated_tokens)]
                    self.history.append(rec)
                    return len(self.history) - 1

    def update_tokens(self, index: int, actual_tokens: int):
        if 0 <= index < len(self.history):
            self.history[index][1] = float(actual_tokens)


class GeminiPool:
    """Пул API-ключей Gemini с равномерной ротацией (Round-Robin)."""
    def __init__(self, keys: List[str], rpm: int = 14, tpm: int = 250000):
        self.keys = keys
        min_interval = max(1.0, 60.0 / rpm + 0.1)
        self.limiters = [KeyRateLimiter(k, rpm=rpm, tpm=tpm, min_interval=min_interval) for k in keys]
        self._current_idx = 0
        self._lock = asyncio.Lock()

    def has_keys(self) -> bool:
        return len(self.keys) > 0

    async def get_next_key_limiter(self) -> Tuple[str, KeyRateLimiter]:
        async with self._lock:
            limiter = self.limiters[self._current_idx]
            self._current_idx = (self._current_idx + 1) % len(self.limiters)
            return limiter.key, limiter


gemini_pool = GeminiPool(
    keys=GEMINI_KEYS,
    rpm=CONFIG["gemini"]["rpm"],
    tpm=CONFIG["gemini"]["tpm"]
)


# ====================== ХЕЛПЕРЫ СЕТИ И PUNYCODE ======================
def to_punycode_url(url: str) -> str:
    """Конвертирует русские домены (.рф) в Punycode для корректной работы DNS."""
    if not url:
        return ""
    try:
        url = url.strip()
        if not re.match(r"^https?://", url, re.IGNORECASE):
            url = f"http://{url}"
        parsed = urlparse(url)
        host = parsed.netloc
        if host:
            port = ""
            if ":" in host:
                host, port = host.split(":", 1)
                port = f":{port}"
            encoded_host = host.encode("idna").decode("ascii")
            netloc = f"{encoded_host}{port}"
            path = quote(parsed.path, safe="/:?#[]@!$&'()*+,;=")
            query = f"?{parsed.query}" if parsed.query else ""
            return f"{parsed.scheme}://{netloc}{path}{query}"
    except Exception:
        pass
    return url


REALISTIC_HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36",
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/avif,image/webp,image/apng,*/*;q=0.8",
    "Accept-Language": "ru-RU,ru;q=0.9,en-US;q=0.8,en;q=0.7",
    "Accept-Encoding": "gzip, deflate, br",
    "Sec-Ch-Ua": '"Chromium";v="124", "Google Chrome";v="124", "Not-A.Brand";v="99"',
    "Sec-Ch-Ua-Mobile": "?0",
    "Sec-Ch-Ua-Platform": '"Windows"',
    "Sec-Fetch-Dest": "document",
    "Sec-Fetch-Mode": "navigate",
    "Sec-Fetch-Site": "none",
    "Sec-Fetch-User": "?1",
    "Upgrade-Insecure-Requests": "1",
}


# ====================== NODE 1: GeoHarvest ======================
async def fetch_2gis_page(client: httpx.AsyncClient, query: str, page: int, api_key: str) -> List[dict]:
    url = "https://catalog.api.2gis.com/3.0/items"
    params = {
        "q": f"{query} {CONFIG['city']}",
        "page": page,
        "page_size": 50,
        "fields": "items.contact_groups,items.point,items.rubrics,items.address,items.org,items.general_info,items.schedule",
        "key": api_key,
        "locale": "ru_RU",
    }
    headers = {
        "User-Agent": REALISTIC_HEADERS["User-Agent"],
        "Referer": "https://2gis.ru/",
    }
    try:
        resp = await client.get(url, params=params, headers=headers, timeout=15)
        if resp.status_code == 200:
            return resp.json().get("result", {}).get("items", [])
    except Exception as e:
        logger.warning(f"2GIS error: {e}")
    return []


async def harvest_2gis() -> List[RawClinic]:
    if not USE_2GIS:
        return []
    logger.info("→ 2GIS сбор...")
    clinics = {}
    async with httpx.AsyncClient(http2=True, verify=False) as client:
        working_key = None
        for key in PUBLIC_KEYS_2GIS:
            if await fetch_2gis_page(client, "косметология", 1, key):
                working_key = key
                logger.info(f"Рабочий ключ 2GIS: {key}")
                break
        if not working_key:
            logger.error("Нет рабочего ключа 2GIS")
            return []

        for query in CONFIG["search_queries"]:
            for page in range(1, CONFIG["max_pages_2gis"] + 1):
                items = await fetch_2gis_page(client, query, page, working_key)
                if not items:
                    break
                for item in items:
                    name = item.get("name", "").strip()
                    if not name or any(sw in name.lower() for sw in CONFIG["stop_words_gov"]):
                        continue
                    sid = str(item.get("id", ""))
                    if sid in clinics:
                        clinics[sid].branches_count += 1
                        continue

                    phones, site = [], ""
                    for g in item.get("contact_groups", []):
                        for c in g.get("contacts", []):
                            if c.get("type") == "phone":
                                phones.append(c.get("value", ""))
                            elif c.get("type") == "website":
                                site = c.get("url") or c.get("value", "")

                    # Извлечение сниппета из описания 2GIS
                    snippet_parts = []
                    if item.get("general_info"):
                        snippet_parts.append(str(item["general_info"]))
                    if item.get("org", {}).get("description"):
                        snippet_parts.append(str(item["org"]["description"]))

                    clinics[sid] = RawClinic(
                        source_id=sid,
                        name=name,
                        address=item.get("address_name") or item.get("full_address_name", ""),
                        phones=phones,
                        site=site,
                        rubrics=[r.get("name", "") for r in item.get("rubrics", [])],
                        source="2gis",
                        snippet=" | ".join(snippet_parts)
                    )
                await asyncio.sleep(0.35)
    logger.info(f"2GIS собрал: {len(clinics)}")
    return list(clinics.values())


async def harvest_yandex() -> List[RawClinic]:
    if not USE_YANDEX:
        return []
    logger.info("→ Яндекс.Карты (Playwright)...")
    try:
        from playwright.async_api import async_playwright
    except ImportError:
        logger.error("Playwright не установлен")
        return []

    clinics = []
    seen = set()

    async with async_playwright() as p:
        browser = await p.chromium.launch(headless=True, args=["--disable-blink-features=AutomationControlled"])
        context = await browser.new_context(
            user_agent=REALISTIC_HEADERS["User-Agent"],
            viewport={"width": 1400, "height": 900}
        )
        page = await context.new_page()

        for query in CONFIG["search_queries"]:
            url = f"https://yandex.ru/maps/75/khabarovsk/search/{quote(query)}"
            logger.info(f"Яндекс: {query}")
            try:
                await page.goto(url, wait_until="domcontentloaded", timeout=35000)
                await page.wait_for_timeout(3500)

                for _ in range(7):
                    await page.mouse.wheel(0, 2200)
                    await page.wait_for_timeout(1100)

                cards = await page.query_selector_all(
                    "div[class*='search-snippet-view'], div[class*='business-snippet'], div[class*='snippet']"
                )
                for card in cards[:CONFIG["max_results_yandex_per_query"]]:
                    try:
                        name_el = await card.query_selector("div[class*='title'] a, a[class*='title'], div[class*='title']")
                        name = (await name_el.inner_text()).strip() if name_el else ""
                        if not name or name.lower() in seen:
                            continue
                        if any(sw in name.lower() for sw in CONFIG["stop_words_gov"]):
                            continue

                        phone = ""
                        phone_el = await card.query_selector("span[class*='phone'], a[href^='tel:']")
                        if phone_el:
                            phone = (await phone_el.inner_text()).strip()

                        site = ""
                        for a in await card.query_selector_all("a[href^='http']"):
                            href = await a.get_attribute("href") or ""
                            if "yandex" not in href and "maps" not in href:
                                site = href
                                break

                        addr_el = await card.query_selector("div[class*='address'], span[class*='address']")
                        address = (await addr_el.inner_text()).strip() if addr_el else ""

                        # Извлечение особенностей/сниппета из карточки Яндекса
                        badge_el = await card.query_selector("div[class*='badge'], div[class*='features'], div[class*='rubric']")
                        snippet = (await badge_el.inner_text()).strip() if badge_el else ""

                        clinics.append(RawClinic(
                            source_id=f"ya_{abs(hash(name + phone))}",
                            name=name,
                            address=address,
                            phones=[phone] if phone else [],
                            site=site,
                            rubrics=[query],
                            source="yandex",
                            snippet=snippet
                        ))
                        seen.add(name.lower())
                    except Exception:
                        continue
            except Exception as e:
                logger.warning(f"Яндекс ({query}): {e}")

        await browser.close()

    logger.info(f"Яндекс собрал: {len(clinics)}")
    return clinics


async def harvest_clinics() -> List[RawClinic]:
    logger.info("Node 1: GeoHarvest (2GIS + Яндекс)")
    tasks = []
    if USE_2GIS:
        tasks.append(harvest_2gis())
    if USE_YANDEX:
        tasks.append(harvest_yandex())

    results = await asyncio.gather(*tasks)
    all_clinics = [c for res in results for c in res]
    logger.info(f"Всего собрано с карт: {len(all_clinics)}")
    return all_clinics


# ====================== УМНАЯ ДЕДУПЛИКАЦИЯ (2GIS + ЯНДЕКС) ======================
def normalize_phone(phone: str) -> str:
    digits = re.sub(r"\D", "", phone)
    if digits.startswith("8") and len(digits) == 11:
        digits = "7" + digits[1:]
    elif digits.startswith("9") and len(digits) == 10:
        digits = "7" + digits
    if len(digits) == 11 and digits.startswith("7"):
        return f"+7 ({digits[1:4]}) {digits[4:7]}-{digits[7:9]}-{digits[9:]}"
    return phone


def phone_digits_key(phone: str) -> str:
    digits = re.sub(r"\D", "", phone)
    if digits.startswith("8") and len(digits) == 11:
        digits = "7" + digits[1:]
    elif digits.startswith("9") and len(digits) == 10:
        digits = "7" + digits
    return digits


def extract_clean_domain(url: str) -> str:
    """Извлекает чистый домен или путь к профилю (VK, Taplink) для сравнения."""
    if not url:
        return ""
    try:
        url_lower = url.lower().strip()
        parsed = urlparse(url_lower if "://" in url_lower else f"http://{url_lower}")
        host = parsed.netloc.lower()
        if host.startswith("www."):
            host = host[4:]
        try:
            host = host.encode('idna').decode('ascii')
        except Exception:
            pass
        if any(soc in host for soc in ["vk.com", "taplink.cc", "t.me"]):
            path = parsed.path.rstrip("/").lower()
            return f"{host}{path}"
        if any(agg in host for agg in ["2gis", "yandex", "zoon", "prodoctorov", "doctu", "yell"]):
            return ""
        return host
    except Exception:
        return ""


def clean_brand_name(name: str) -> str:
    s = name.lower()
    s = re.sub(r"[«»\"'“”„`]", " ", s)
    s = re.sub(r"\b(ооо|зао|пао|ао|ип|клиника|центр|сеть|салон|красоты|медицинский|медицинская|стоматология|стоматологическая|косметология|косметологическая|кабинет|доктор|врач|отделение)\b", " ", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s


def extract_street_keyword(address: str) -> str:
    if not address:
        return ""
    addr_l = address.lower()
    addr_l = re.sub(r"\b(г\.|хабаровск|ул\.|улица|пр-кт|проспект|пер\.|переулок|д\.|дом|к\.|корпус|оф\.|офис|литер|эт\.|этаж)\b", " ", addr_l)
    tokens = [t for t in re.split(r"[\s,]+", addr_l) if len(t) > 3 and not t.isdigit()]
    return tokens[0] if tokens else ""


def smart_cascade_dedup(clinics: List[RawClinic]) -> List[RawClinic]:
    """Умное каскадное слияние дублей 2GIS и Яндекс по доменам, телефонам и бренду."""
    logger.info(f"Node 2: Умное каскадное объединение ({len(clinics)} сырых записей)...")
    unique_clinics: List[RawClinic] = []

    domain_index: Dict[str, int] = {}
    phone_index: Dict[str, int] = {}
    brand_street_index: Dict[Tuple[str, str], int] = {}
    exact_name_index: Dict[str, int] = {}

    for c in clinics:
        # Нормализация телефонов
        c.phones = [normalize_phone(p) for p in c.phones if p]

        match_idx: Optional[int] = None

        # 1. Поиск по чистому домену / профилю VK / Taplink
        dom = extract_clean_domain(c.site)
        if dom and dom in domain_index:
            match_idx = domain_index[dom]

        # 2. Поиск по номерам телефонов
        if match_idx is None:
            for p in c.phones:
                pkey = phone_digits_key(p)
                if len(pkey) >= 10 and pkey in phone_index:
                    match_idx = phone_index[pkey]
                    break

        # 3. Поиск по бренду и улице
        brand = clean_brand_name(c.name)
        street = extract_street_keyword(c.address)
        if match_idx is None and len(brand) >= 3 and street:
            key = (brand, street)
            if key in brand_street_index:
                match_idx = brand_street_index[key]

        # 4. Поиск по точному имени
        norm_name = re.sub(r"\s+", " ", c.name.lower().strip())
        if match_idx is None and norm_name in exact_name_index:
            match_idx = exact_name_index[norm_name]

        if match_idx is not None:
            # Слияние с существующей записью
            existing = unique_clinics[match_idx]
            if existing.source != c.source:
                existing.source = "2gis + yandex"
            # Объединяем телефоны
            for p in c.phones:
                if p not in existing.phones:
                    existing.phones.append(p)
            # Объединяем рубрики
            for r in c.rubrics:
                if r not in existing.rubrics:
                    existing.rubrics.append(r)
            # Выбираем лучший сайт (отдельный сайт > VK/Taplink > пусто)
            if not existing.site and c.site:
                existing.site = c.site
            elif existing.site and c.site and "vk.com" in existing.site and "vk.com" not in c.site:
                existing.site = c.site
            # Выбираем более полный адрес
            if len(c.address) > len(existing.address):
                existing.address = c.address
            # Объединяем сниппеты
            if c.snippet and c.snippet not in existing.snippet:
                existing.snippet = (existing.snippet + " | " + c.snippet).strip(" | ")
            existing.branches_count = max(existing.branches_count, c.branches_count)
        else:
            new_idx = len(unique_clinics)
            unique_clinics.append(c)

            if dom:
                domain_index[dom] = new_idx
            for p in c.phones:
                pkey = phone_digits_key(p)
                if len(pkey) >= 10:
                    phone_index[pkey] = new_idx
            if len(brand) >= 3 and street:
                brand_street_index[(brand, street)] = new_idx
            if norm_name:
                exact_name_index[norm_name] = new_idx

    logger.info(f"Уникальных после умного слияния: {len(unique_clinics)} (исключено {len(clinics) - len(unique_clinics)} дублей)")
    return unique_clinics


# ====================== NODE 3: SmartCrawler ======================
async def fetch_with_httpx(client: httpx.AsyncClient, url: str) -> Tuple[str, str]:
    target_url = to_punycode_url(url)
    cache_key = f"httpx:{target_url}"
    if cache_key in cache:
        return cache[cache_key], "cached"
    try:
        headers = dict(REALISTIC_HEADERS)
        headers["Referer"] = target_url
        resp = await client.get(target_url, headers=headers, follow_redirects=True, timeout=CONFIG["crawler"]["timeout"])
        if resp.status_code == 200 and len(resp.text) > 500:
            cache[cache_key] = resp.text
            return resp.text, "ok"
        return "", f"http_{resp.status_code}"
    except Exception as e:
        return "", f"error_{type(e).__name__}"


async def fetch_with_playwright(url: str) -> Tuple[str, str]:
    target_url = to_punycode_url(url)
    cache_key = f"pw:{target_url}"
    if cache_key in cache:
        return cache[cache_key], "cached_pw"
    try:
        from playwright.async_api import async_playwright
        async with async_playwright() as p:
            browser = await p.chromium.launch(headless=True, args=["--disable-blink-features=AutomationControlled"])
            context = await browser.new_context(user_agent=REALISTIC_HEADERS["User-Agent"], ignore_https_errors=True)
            page = await context.new_page()
            await page.goto(target_url, wait_until="domcontentloaded", timeout=20000)
            await page.wait_for_timeout(2500)
            html = await page.content()
            await browser.close()
            if len(html) > 500:
                cache[cache_key] = html
                return html, "ok_pw"
            return "", "empty_pw"
    except Exception as e:
        logger.warning(f"Playwright {url}: {e}")
        return "", "error_pw"


def extract_links(html: str, base_url: str) -> List[str]:
    soup = BeautifulSoup(html, "lxml")
    patterns = [
        r"(price|prays|uslug|service|tsen|cost|цен|прайс|услуг)",
        r"(vrach|doctor|specialist|team|kollektiv|врач|специалист|команда|о-нас)",
        r"(contact|kontakty|контакт)"
    ]
    links = []
    for a in soup.find_all("a", href=True):
        href = a.get("href", "").lower()
        text = a.get_text(" ", strip=True).lower()
        if any(re.search(p, href + " " + text) for p in patterns):
            full = urljoin(base_url, a["href"])
            if full.startswith("http"):
                links.append(full)
    return list(dict.fromkeys(links))[:CONFIG["crawler"]["max_pages_per_site"]]


async def crawl_one(clinic: RawClinic, client: httpx.AsyncClient) -> Dict[str, Any]:
    # 1. Если сайта нет вообще — сразу используем метаданные карточки
    if not clinic.site:
        card_text = f"Организация: {clinic.name}. Адрес: {clinic.address}. Рубрики и деятельность: {', '.join(clinic.rubrics)}. {clinic.snippet}".strip()
        return {"text": card_text, "emails": [], "status": "no_site_card"}

    html, status = await fetch_with_httpx(client, clinic.site)

    # 2. Если 403 / 503 / мало данных — пробуем Playwright
    if (not html or "http_403" in status or "http_503" in status or len(html) < 800) and CONFIG["crawler"]["use_playwright_fallback"]:
        html, status = await fetch_with_playwright(clinic.site)

    # 3. ФОЛЛБЕК НА КАРТОЧКУ: если сайт заблокирован или недоступен
    if not html:
        card_text = (
            f"Организация: {clinic.name}. Сайт (недоступен): {clinic.site}. "
            f"Адрес: {clinic.address}. Рубрики и деятельность: {', '.join(clinic.rubrics)}. "
            f"Описание и сниппет: {clinic.snippet}"
        ).strip()
        return {"text": card_text, "emails": [], "status": f"card_fallback_{status}"}

    # Парсинг успешной страницы
    soup = BeautifulSoup(html, "lxml")
    all_text = soup.get_text(separator=" ", strip=True)

    emails = set(re.findall(r"[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}", all_text))
    for a in soup.select('a[href^="mailto:"]'):
        emails.add(a["href"].replace("mailto:", "").split("?")[0])

    # Обход подстраниц
    for link in extract_links(html, clinic.site)[:3]:
        page_html, _ = await fetch_with_httpx(client, link)
        if page_html:
            all_text += " " + BeautifulSoup(page_html, "lxml").get_text(" ", strip=True)

    return {"text": all_text[:35000], "emails": list(emails)[:3], "status": status}


async def crawl_sites(clinics: List[RawClinic]) -> List[Dict]:
    logger.info(f"Node 3: SmartCrawler для {len(clinics)} организаций (с обходом SSL и Punycode)...")
    semaphore = asyncio.Semaphore(CONFIG["crawler"]["concurrency"])
    async with httpx.AsyncClient(http2=True, verify=False, timeout=30) as client:
        async def worker(c):
            async with semaphore:
                data = await crawl_one(c, client)
                return {**asdict(c), **data}
        return await asyncio.gather(*(worker(c) for c in clinics))


# ====================== NODE 4: HybridAnalyzer ======================
def regex_analysis(text: str) -> Dict[str, str]:
    text_l = text.lower()
    inj_patterns = (
        r"(биоревитализ|мезотерапи|ботулотоксин|ботокс|диспорт|ксеомин|релатокс|миотокс|"
        r"контурная\s+пластик|филлер|мезонит|нитевой\s+лифтинг|коллагенотерапи|сферогель|"
        r"коллост|nithya|нития|полимолочн|aesthefill|sculptra|miraline|"
        r"плазмолифтинг\s+десен|prp\s+десен|инъекци.*?десн|гиалуронов.*?десн|бруксизм|"
        r"внутрисуставн.*?инъекци|лечебн.*?блокад|гиалуронов.*?сустав|протез\s+синовиальн|"
        r"интимн.*?филлинг|интимн.*?биоревитализ|эстетическ.*?гинеколог)"
    )
    inj = bool(re.search(inj_patterns, text_l))

    prp = bool(re.search(r"(плазмолифтинг|плазмотерапи|аутоплазм|prp(\s*-\s*терапия)?|prgf)", text_l))
    brands = [b for b, vars in CONFIG["prp_brands"].items() if any(v in text_l for v in vars)]
    prp_str = "Нет"
    if prp:
        prp_str = f"Да ({', '.join(brands)})" if brands else "Да (бренд не указан)"

    return {"injection_therapy": "Да" if inj else "Нет", "prp_therapy_brand": prp_str}


async def call_gemini(site_text: str) -> Dict[str, Any]:
    if not gemini_pool.has_keys():
        return {}

    prompt = GEMINI_PROMPT.format(site_text=site_text[:28000])
    estimated_tokens = max(200, int(len(prompt) / 2.5) + 300)

    max_retries = 3
    backoff = 5.0

    for attempt in range(1, max_retries + 1):
        key, limiter = await gemini_pool.get_next_key_limiter()
        slot_idx = await limiter.acquire(estimated_tokens)
        try:
            import google.generativeai as genai
            genai.configure(api_key=key)
            model = genai.GenerativeModel(CONFIG["gemini"]["model"])
            response = await asyncio.to_thread(
                model.generate_content, prompt,
                generation_config={"temperature": CONFIG["gemini"]["temperature"], "response_mime_type": "application/json"}
            )
            if hasattr(response, "usage_metadata") and response.usage_metadata:
                total_tokens = getattr(response.usage_metadata, "total_token_count", None)
                if total_tokens:
                    limiter.update_tokens(slot_idx, total_tokens)

            return json.loads(response.text)
        except Exception as e:
            err_msg = str(e).lower()
            if any(term in err_msg for term in ("429", "resource_exhausted", "quota", "rate limit")):
                if attempt < max_retries:
                    logger.warning(f"Gemini квота/пауза (попытка {attempt}/{max_retries}): {e}. Пауза {backoff}с...")
                    await asyncio.sleep(backoff)
                    backoff *= 2
                    continue
            logger.warning(f"Gemini error: {e}")
            return {}

    return {}


# ====================== ЧЕКПОИНТЫ И АНАЛИЗ ======================
def load_checkpoint() -> Dict[str, dict]:
    if CHECKPOINT_FILE.exists():
        try:
            with open(CHECKPOINT_FILE, "r", encoding="utf-8") as f:
                data = json.load(f)
                logger.info(f"Загружен чекпоинт: {len(data)} клиник уже проанализировано")
                return data
        except Exception as e:
            logger.warning(f"Ошибка загрузки чекпоинта: {e}")
    return {}


def save_checkpoint(checkpoint: Dict[str, dict]):
    try:
        with open(CHECKPOINT_FILE, "w", encoding="utf-8") as f:
            json.dump(checkpoint, f, ensure_ascii=False, indent=2)
    except Exception as e:
        logger.warning(f"Ошибка сохранения чекпоинта: {e}")


async def analyze_clinics(crawled: List[Dict]) -> List[ClinicRecord]:
    logger.info("Node 4: HybridAnalyzer (анализ сайтов и карточек)...")
    checkpoint = load_checkpoint()
    records: List[ClinicRecord] = []
    total = len(crawled)

    def sigint_handler(signum, frame):
        logger.warning(f"\n[Ctrl+C] Прерывание процесса! Экстренное сохранение {len(records)} записей...")
        save_checkpoint(checkpoint)
        export_to_excel(records)
        sys.exit(0)

    original_handler = signal.getsignal(signal.SIGINT)
    signal.signal(signal.SIGINT, sigint_handler)

    try:
        for i, item in enumerate(crawled, 1):
            sid = item.get("source_id", f"c_{i}")

            # Пропуск уже обработанных
            if sid in checkpoint:
                rec_dict = checkpoint[sid]
                rec_dict["№ п/п"] = i
                rec = ClinicRecord(**rec_dict)
                records.append(rec)
                continue

            text = item.get("text", "")
            status = item.get("status", "ok")
            reg = regex_analysis(text)
            injection = reg["injection_therapy"]
            prp_brand = reg["prp_therapy_brand"]
            specialists = "Требует уточнения"
            note = ""

            # Если контент получен (сайт или карточка) — запускаем Gemini
            if text and len(text) > 30:
                gem = await call_gemini(text)
                if gem:
                    if gem.get("injection_therapy") in ("Да", "Нет"):
                        injection = gem["injection_therapy"]
                    if gem.get("prp_therapy") == "Да":
                        brands = gem.get("prp_brands") or []
                        prp_brand = f"Да ({', '.join(brands)})" if brands else "Да (бренд не указан)"
                    elif gem.get("prp_therapy") == "Нет":
                        prp_brand = "Нет"
                    sc = gem.get("specialists_count")
                    specialists = str(sc) if sc is not None else "Требует уточнения"
                    note = str(gem.get("specialists_comment", ""))[:150]

            if "card_fallback" in status:
                note = f"По карточке (сайт не ответил): {note}" if note else "По карточке (сайт не ответил)"
            elif status == "no_site_card":
                note = f"По карточке (без сайта): {note}" if note else "По карточке (без сайта)"

            record = ClinicRecord(
                index=i,
                name=item["name"],
                branches_count=item.get("branches_count", 1),
                activity=", ".join(item.get("rubrics", [])),
                phone=item["phones"][0] if item.get("phones") else "",
                email=(item.get("emails") or [""])[0],
                site=item.get("site", ""),
                specialists_count=specialists,
                injection_therapy=injection,
                prp_therapy_brand=prp_brand,
                status=status,
                note=note,
                source=item.get("source", ""),
                source_id=sid
            )
            records.append(record)
            checkpoint[sid] = record.model_dump(by_alias=True)

            logger.info(
                f"[{i}/{total}] ({i/total*100:.1f}%) {record.name[:30]:<30} | "
                f"Инъекции: {injection:<5} | PRP: {prp_brand}"
            )

            if i % 15 == 0:
                save_checkpoint(checkpoint)
                export_to_excel(records)

    finally:
        signal.signal(signal.SIGINT, original_handler)
        save_checkpoint(checkpoint)

    return records


# ====================== NODE 5: Exporter ======================
def format_worksheet(ws, df, is_target: bool = False):
    header_fill = PatternFill(
        start_color="1F4E78" if is_target else "404040",
        end_color="1F4E78" if is_target else "404040",
        fill_type="solid"
    )
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    cell_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    for col_num, column_title in enumerate(df.columns, 1):
        cell = ws.cell(row=1, column=col_num, value=column_title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_num, row_data in enumerate(df.values, 2):
        for col_num, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_num, column=col_num, value=val)
            cell.font = cell_font
            cell.border = thin_border
            align = "center" if col_num in (1, 3, 8, 9) else "left"
            cell.alignment = Alignment(horizontal=align, vertical="center")

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col_letter].width = max(min(max_len + 4, 45), 12)

    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"


def export_to_excel(records: List[ClinicRecord], path: str = OUTPUT_EXCEL):
    if not records:
        return

    Path(path).parent.mkdir(parents=True, exist_ok=True)

    target_columns = [
        "№ п/п",
        "Название",
        "кол-во филиалов",
        "Деятельность",
        "Контакт номер",
        "Email",
        "Сайт",
        "Кол-во спецов",
        "Инъекционная терапия",
        "Плазмотерапия (бренд)",
    ]

    raw_dicts = [r.model_dump(by_alias=True) for r in records]
    full_df = pd.DataFrame(raw_dicts)
    full_df = full_df[[col for col in target_columns if col in full_df.columns]]
    full_df["№ п/п"] = range(1, len(full_df) + 1)

    target_df = full_df[full_df["Инъекционная терапия"] == "Да"].copy()
    target_df["№ п/п"] = range(1, len(target_df) + 1)

    def _write_book(target_file: str):
        with pd.ExcelWriter(target_file, engine="openpyxl") as writer:
            wb = writer.book
            ws_target = wb.create_sheet(title="Целевые клиники (Инъекции Да)")
            format_worksheet(ws_target, target_df, is_target=True)
            ws_all = wb.create_sheet(title="Все организации (Архив)")
            format_worksheet(ws_all, full_df, is_target=False)
            if "Sheet" in wb.sheetnames:
                del wb["Sheet"]

    saved_path = path
    try:
        _write_book(saved_path)
    except PermissionError:
        saved_path = path.replace(".xlsx", f"_{int(time.time())}.xlsx")
        logger.warning(f"Файл {path} открыт в Excel. Сохранено в резервную копию: {saved_path}")
        _write_book(saved_path)

    logger.info(
        f"Таблица обновлена: {saved_path} | "
        f"Целевых с инъекциями: {len(target_df)} из {len(full_df)}"
    )


# ====================== MAIN ======================
async def main():
    start = time.time()
    logger.info("=" * 65)
    logger.info("ЗАПУСК ПАЙПЛАЙНА АНАЛИЗА КЛИНИК (ХАБАРОВСК)")
    logger.info(f"Пул Gemini ключей: {len(gemini_pool.keys)} шт. | Модель: {CONFIG['gemini']['model']}")
    logger.info("=" * 65)

    force_reharvest = os.getenv("FORCE_REHARVEST", "false").lower() in ("true", "1", "yes")
    force_recrawl = os.getenv("FORCE_RECRAWL", "false").lower() in ("true", "1", "yes")

    # 1. Загрузка или сбор гео-данных с умной дедупликацией
    if RAW_FILE.exists() and not force_reharvest:
        logger.info(f"Используем кэшированные гео-данные из {RAW_FILE}")
        with open(RAW_FILE, "r", encoding="utf-8") as f:
            cleaned = [RawClinic(**item) for item in json.load(f)]
        cleaned = smart_cascade_dedup(cleaned)
    else:
        raw = await harvest_clinics()
        cleaned = smart_cascade_dedup(raw)
        with open(RAW_FILE, "w", encoding="utf-8") as f:
            json.dump([asdict(c) for c in cleaned], f, ensure_ascii=False, indent=2)
        logger.info(f"Гео-данные сохранены в {RAW_FILE}")

    # 2. Загрузка или краулинг сайтов
    if CRAWLED_FILE.exists() and not force_recrawl:
        logger.info(f"Используем кэшированные сайты из {CRAWLED_FILE}")
        with open(CRAWLED_FILE, "r", encoding="utf-8") as f:
            crawled = json.load(f)
    else:
        crawled = await crawl_sites(cleaned)
        with open(CRAWLED_FILE, "w", encoding="utf-8") as f:
            json.dump(crawled, f, ensure_ascii=False, indent=2)
        logger.info(f"Сайты сохранены в {CRAWLED_FILE}")

    # 3. Анализ клиник (с поддержкой докачки)
    analyzed = await analyze_clinics(crawled)

    # 4. Финальный экспорт
    export_to_excel(analyzed)

    target_count = sum(1 for r in analyzed if r.injection_therapy == "Да")
    logger.info("=" * 65)
    logger.info(f"ГОТОВО за {time.time() - start:.1f} сек.")
    logger.info(f"Всего клиник в базе: {len(analyzed)}")
    logger.info(f"Целевых с инъекционной терапией: {target_count}")
    logger.info(f"Файл сохранен: {OUTPUT_EXCEL}")
    logger.info("=" * 65)


if __name__ == "__main__":
    asyncio.run(main())