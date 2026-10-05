"""Форматированный экспорт результатов анализа в Excel (openpyxl)."""

import time
from pathlib import Path
from typing import List
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from clinic_scraper.core.logger import logger
from clinic_scraper.models import ClinicRecord


def format_worksheet(ws, df: pd.DataFrame, is_target: bool = False) -> None:
    """Применяет корпоративный стиль оформления листа Excel."""
    header_fill = PatternFill(
        start_color="1F4E78" if is_target else "404040",
        end_color="1F4E78" if is_target else "404040",
        fill_type="solid"
    )
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    cell_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    for col_num, column_title in enumerate(df.columns, 1):
        cell = ws.cell(row=1, column=col_num, value=column_title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_num, row_data in enumerate(df.values, 2):
        for col_num, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_num, column=col_num, value=val)
            cell.font = cell_font
            cell.border = thin_border
            align = "center" if col_num in (1, 3, 8, 9) else "left"
            cell.alignment = Alignment(horizontal=align, vertical="center")

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col_letter].width = max(min(max_len + 4, 45), 12)

    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"


def export_to_excel(records: List[ClinicRecord], output_path: Path) -> Path:
    """Генерирует двухстраничный отчет Excel."""
    if not records:
        logger.warning("Нет записей для экспорта в Excel")
        return output_path

    output_path.parent.mkdir(parents=True, exist_ok=True)

    target_columns = [
        "№ п/п",
        "Название",
        "кол-во филиалов",
        "Деятельность",
        "Контакт номер",
        "Email",
        "Сайт",
        "Кол-во спецов",
        "Инъекционная терапия",
        "Плазмотерапия (бренд)",
    ]

    raw_dicts = [r.model_dump(by_alias=True) for r in records]
    full_df = pd.DataFrame(raw_dicts)
    full_df = full_df[[col for col in target_columns if col in full_df.columns]]
    full_df["№ п/п"] = range(1, len(full_df) + 1)

    target_df = full_df[full_df["Инъекционная терапия"] == "Да"].copy()
    target_df["№ п/п"] = range(1, len(target_df) + 1)

    def _write_book(target_file: Path) -> None:
        with pd.ExcelWriter(target_file, engine="openpyxl") as writer:
            wb = writer.book
            ws_target = wb.create_sheet(title="Целевые клиники (Инъекции Да)")
            format_worksheet(ws_target, target_df, is_target=True)
            ws_all = wb.create_sheet(title="Все организации (Архив)")
            format_worksheet(ws_all, full_df, is_target=False)
            if "Sheet" in wb.sheetnames:
                del wb["Sheet"]

    saved_path = output_path
    try:
        _write_book(saved_path)
    except PermissionError:
        saved_path = output_path.with_name(f"{output_path.stem}_{int(time.time())}.xlsx")
        logger.warning(f"Файл {output_path} занят другим процессом. Сохраняем в: {saved_path}")
        _write_book(saved_path)

    logger.info(
        f"Таблица Excel сохранена: {saved_path} | "
        f"Целевых с инъекциями: {len(target_df)} из {len(full_df)}"
    )
    return saved_path
